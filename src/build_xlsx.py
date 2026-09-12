import pickle, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
ff = pickle.load(open('final_fit.pkl','rb')); hr = pickle.load(open('hourly_recon.pkl','rb')); LO = pickle.load(open('loo.pkl','rb'))
E = pd.read_pickle('exp_results.pkl'); d = pd.read_pickle('daily.pkl'); qc = pd.read_csv('qc_log.csv', index_col=0).iloc[:,0]
gaps = pd.read_csv('daily_gaps.csv')
F = 'Arial'; HF = Font(name=F, bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='1F4E79')
BF = Font(name=F); BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=13)
wb = Workbook()
def sheet(name, df, widths=None, numfmt=None, first=False):
    ws = wb.active if first else wb.create_sheet(name); ws.title = name
    ws.append(list(df.columns))
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v,float) and np.isnan(v)) else (v.to_pydatetime() if isinstance(v,pd.Timestamp) else v) for v in row])
    for j, col in enumerate(df.columns, 1):
        L = get_column_letter(j); ws.column_dimensions[L].width = (widths or {}).get(col, max(12, min(40, len(str(col))+2)))
        if numfmt and col in numfmt:
            for cell in ws[L][1:]: cell.number_format = numfmt[col]
    ws.freeze_panes = 'A2'
    return ws
# ---------------- README ----------------
ws = wb.active; ws.title = 'README'
lines = [
 ('Bang Pakong water level: Gaussian-process gap reconstruction (application of Xu et al., 2024)', TITLE),
 ('Source data: Clean_BangPaKong 2.xlsx (user-supplied), 353,644 records, 7 May 2019 to 23 Jun 2026, nominal 10-min sampling.', BF),
 ('Method reference: Xu K. et al. (2024) Reconstruction of geodetic time series with missing data and time-varying seasonal signals using Gaussian process for machine learning. GPS Solutions 28:79. doi:10.1007/s10291-024-01616-8', BF),
 ('', BF),
 ('Processing chain', BOLD),
 ('1. Quality control of 10-min data (see QC_Log): duplicate timestamps merged, zero readings, stuck high readings (>= 5.60 m), flat-line runs >= 3 h and Hampel-type spikes (> 0.25 m from a 70-min running median) removed.', BF),
 ('2. Hourly series: on-the-hour values by linear interpolation between neighbouring readings no more than 30 min apart.', BF),
 ('3. Tidal harmonic analysis: 33 constituents (diurnal to M8) with nodal corrections, least squares on all hourly data (see Tidal_Constituents).', BF),
 ('4. Daily mean water level (DMWL): daily mean of the non-tidal residual plus Z0; days with fewer than 18 valid hours set to missing.', BF),
 ('5. Models for DMWL: Standard (linear trend + constant annual and semiannual terms); Bennett-type time-varying seasonal model (Kalman smoother, random-walk seasonal coefficients + AR(1) noise); GP with mean h(t)=[1, t, sin 2πt, cos 2πt] and Matérn 3/2 kernel (Xu et al. 2024 set-up); GP-2M-SA extension (two Matérn 3/2 components, annual + semiannual mean). Hyperparameters by maximum likelihood.', BF),
 ('6. Gap-recovery experiments: artificial gaps (contiguous interior, end-of-record, realistic Jun–Sep and Nov–Jan windows, scattered 1–10-day gaps); models refitted on the remaining data each time.', BF),
 ('7. Hourly reconstruction: hourly value = tidal prediction + GP-2M-SA daily level (interpolated to hours). Gaps <= 6 h also receive a linearly interpolated non-tidal residual.', BF),
 ('', BF),
 ('Sheets', BOLD),
 ('Daily_Reconstructed – observed and modelled DMWL on a continuous daily grid, 95% intervals (formula columns), forecast to 31 Dec 2026.', BF),
 ('Hourly_GapFilled – continuous hourly series with tide, low-frequency level, reconstruction, uncertainty and flag (values).', BF),
 ('Model_Fit – full-record model comparison: rate, uncertainty, likelihood, AIC, in-sample and leave-one-out RMSE, hyperparameters.', BF),
 ('Gap_Experiments / Experiment_Summary – per-case metrics and AVERAGEIFS summaries by scenario and model.', BF),
 ('Hourly_Validation / Hourly_Val_Summary – hourly skill of reconstruction methods on 10 withheld windows.', BF),
 ('Tidal_Constituents, Daily_Gaps, QC_Log – supporting tables.', BF),
 ('', BF),
 ('Flags', BOLD),
 ('Daily flag: 0 = observed; 1 = reconstructed (missing day); 2 = forecast beyond the record.', BF),
 ('Hourly flag: 0 = observed; 1 = short gap (<= 6 h) filled with tide + level + interpolated residual; 2 = model-filled (tide + GP-2M-SA level).', BF),
 ('Units: metres (water level), cm (error metrics), mm/yr (rates). Datum as in the source file. Time as in the source file (assumed local time).', BF),
 ('Notes: GP in-sample RMSE is near zero because the fitted noise variance is very small; use the leave-one-out and gap-experiment RMSE for model comparison.', BF),
]
for i,(txt,fnt) in enumerate(lines,1):
    c = ws.cell(row=i, column=1, value=txt); c.font = fnt; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 140
# ---------------- QC log ----------------
qdesc = {'raw_records':'Records in source file','duplicate_timestamps':'Timestamps with more than one record (merged to median)',
 'conflicting_duplicates_removed':'Timestamps whose duplicates disagree by > 0.10 m (removed)','zero_values_removed':'Readings equal to 0 m',
 'stuck_high_values_removed(>=5.60 m)':'Stuck/saturated readings >= 5.60 m (5.67–5.69 m plateau, Nov–Dec 2022)',
 'flatline_values_removed(run>=3h)':'Identical consecutive readings lasting >= 3 h','spikes_removed_pass1':'Spikes > 0.25 m from 70-min running median (pass 1)',
 'spikes_removed_pass2':'Spikes, pass 2','clean_records':'Unique timestamps retained','hourly_grid_points':'Hourly grid length',
 'hourly_missing':'Hourly values missing after QC','hourly_missing_pct':'Hourly values missing (%)'}
sheet('QC_Log', pd.DataFrame({'step':qc.index, 'description':[qdesc.get(k,'') for k in qc.index], 'count':qc.values}), widths={'step':34,'description':70,'count':12})
# ---------------- Daily ----------------
gi = ff['grid_idx']; n = len(d)
dm = pd.DataFrame({'date':gi})
dm['hours_valid'] = list(d.hours.values) + [None]*(len(gi)-n)
dm['observed_dmwl_m'] = list(d.dmwl.values) + [np.nan]*(len(gi)-n)
dm['St_m'] = ff['St']['mean']
dm['TVS_m'] = list(ff['TVS']['mean']) + [np.nan]*(len(gi)-n)
dm['GP_m'] = ff['GP']['mean']; dm['GP_sd_m'] = ff['GP']['sd']
dm['GP2MSA_m'] = ff['GP2M32SA']['mean']; dm['GP2MSA_sd_m'] = ff['GP2M32SA']['sd']
dm['GP2MSA_lower95_m'] = None; dm['GP2MSA_upper95_m'] = None
dm['final_dmwl_m'] = None
dm['flag'] = np.where(np.arange(len(gi))>=n, 2, np.where(np.isnan(dm.observed_dmwl_m), 1, 0))
ws = sheet('Daily_Reconstructed', dm, widths={'date':12}, numfmt={'date':'yyyy-mm-dd', **{c:'0.000' for c in ['observed_dmwl_m','St_m','TVS_m','GP_m','GP_sd_m','GP2MSA_m','GP2MSA_sd_m','GP2MSA_lower95_m','GP2MSA_upper95_m','final_dmwl_m']}})
for r in range(2, len(gi)+2):
    ws[f'J{r}'] = f'=H{r}-1.96*I{r}'; ws[f'K{r}'] = f'=H{r}+1.96*I{r}'
    ws[f'L{r}'] = f'=IF(ISNUMBER(C{r}),C{r},H{r})'
    for c in 'JKL': ws[f'{c}{r}'].number_format = '0.000'
# ---------------- Hourly ----------------
H = hr['hourly'].reset_index()
H['lower95_m'] = H.reconstructed_m - 1.96*H.sd_m; H['upper95_m'] = H.reconstructed_m + 1.96*H.sd_m
sheet('Hourly_GapFilled', H, widths={'datetime':18}, numfmt={'datetime':'yyyy-mm-dd hh:mm', **{c:'0.000' for c in ['observed_m','tide_m','low_freq_level_m','reconstructed_m','sd_m','lower95_m','upper95_m']}})
# ---------------- Model fit ----------------
T = ff['table'].copy()
loo_map = {0:LO['rmse']['St'], 1:np.nan, 2:LO['rmse']['GP'], 3:LO['rmse']['GP2M32'], 4:LO['rmse']['GPM12'], 5:LO['rmse']['GPQP'], 6:LO['rmse']['GP2M32SA']}
T['LOO_RMSE_cm'] = [loo_map[i] for i in range(len(T))]
T.loc[1,'model_RMSE_cm'] = 100*np.sqrt(np.nanmean((d.dmwl.values - ff['TVS']['det'])**2))
T = T.rename(columns={'model_RMSE_cm':'in_sample_RMSE_cm','rate_sd_mm_yr':'rate_1sd_mm_yr','rate_sd_whitenoise_mm_yr':'rate_1sd_white_noise_only_mm_yr','acf1':'lag1_acf_in_sample_resid'})
keep = ['model','rate_mm_yr','rate_1sd_mm_yr','rate_1sd_white_noise_only_mm_yr','in_sample_RMSE_cm','LOO_RMSE_cm','nll','AIC','n_hyper'] + [c for c in T.columns if c.startswith('hp_')]
T = T[keep]
ws = sheet('Model_Fit', T, widths={'model':58}, numfmt={c:'0.00' for c in ['rate_mm_yr','rate_1sd_mm_yr','rate_1sd_white_noise_only_mm_yr','in_sample_RMSE_cm','LOO_RMSE_cm']} | {'nll':'0.0','AIC':'0.0'} | {c:'0.000000' for c in T.columns if c.startswith('hp_')})
r = len(T)+3
notes = ['Notes: rates are relative water-level trends from the daily series (not corrected for vertical land motion).',
 'St rate uncertainty uses a first-order Gauss–Markov + white-noise covariance (MLE); the white-noise-only value is shown for comparison.',
 'TVS in-sample RMSE uses the trend + seasonal part only; its likelihood is not directly comparable with the GP likelihoods (first 60 innovations excluded).',
 'Kernel length scales (ell) are in years: 0.0023 yr ≈ 0.9 day; 0.0154 yr ≈ 5.6 days; 0.0049 yr ≈ 1.8 days.',
 'LOO_RMSE = closed-form leave-one-out residuals with hyperparameters fixed at the full-record MLE.']
for i,t_ in enumerate(notes): ws.cell(row=r+i, column=1, value=t_).font = Font(name=F, italic=True)
# ---------------- Experiments ----------------
Ex = E[['case','group','label','pct','model','n_train','n_test','RMSE_cm','MAE_cm','bias_cm','cover95_pct','rate_mm_yr']].sort_values(['case','model'])
sheet('Gap_Experiments', Ex, widths={'group':26,'label':30}, numfmt={c:'0.00' for c in ['RMSE_cm','MAE_cm','bias_cm','cover95_pct','rate_mm_yr']})
ws = wb.create_sheet('Experiment_Summary')
models = ['LI','St','TVS','GP','GP-2M','GP-2M-SA']; groups = ['A_interior_block','B_end_prediction','C_realistic_JunSep_75d','C_realistic_NovJan_39d','D_scattered_1to10d']
N = len(Ex)+1
ws['A1'] = 'Mean RMSE (cm) by scenario and model — AVERAGEIFS over Gap_Experiments'; ws['A1'].font = BOLD
ws.append(['scenario']+models)
for c in ws[2]: c.font = HF; c.fill = HFILL
for gi_, g in enumerate(groups):
    row = 3+gi_; ws.cell(row=row, column=1, value=g).font = BF
    for j, m in enumerate(models, 2):
        c = ws.cell(row=row, column=j, value=f'=AVERAGEIFS(Gap_Experiments!$H$2:$H${N},Gap_Experiments!$B$2:$B${N},$A{row},Gap_Experiments!$E$2:$E${N},{get_column_letter(j)}$2)')
        c.number_format = '0.00'; c.font = BF
row = 3+len(groups); ws.cell(row=row, column=1, value='All scenarios').font = BOLD
for j, m in enumerate(models, 2):
    c = ws.cell(row=row, column=j, value=f'=AVERAGEIFS(Gap_Experiments!$H$2:$H${N},Gap_Experiments!$E$2:$E${N},{get_column_letter(j)}$2)'); c.number_format='0.00'; c.font=BOLD
r2 = row+2
ws.cell(row=r2, column=1, value='Mean 95% interval coverage (%)').font = BOLD
ws.cell(row=r2+1, column=1, value='scenario')
for j, m in enumerate(['TVS','GP','GP-2M','GP-2M-SA'], 2): ws.cell(row=r2+1, column=j, value=m)
for c in ws[r2+1]: 
    if c.value: c.font = HF; c.fill = HFILL
for gi_, g in enumerate(groups):
    rr = r2+2+gi_; ws.cell(row=rr, column=1, value=g)
    for j, m in enumerate(['TVS','GP','GP-2M','GP-2M-SA'], 2):
        c = ws.cell(row=rr, column=j, value=f'=AVERAGEIFS(Gap_Experiments!$K$2:$K${N},Gap_Experiments!$B$2:$B${N},$A{rr},Gap_Experiments!$E$2:$E${N},{get_column_letter(j)}${r2+1})'); c.number_format='0.0'
ws.cell(row=r2+8, column=1, value='LI = linear interpolation (persistence of the last value for end-of-record prediction). St = Standard; TVS = Bennett-type Kalman model; GP = Matérn 3/2 (Xu et al. 2024); GP-2M = two Matérn 3/2 components; GP-2M-SA = GP-2M with annual + semiannual mean.').font = Font(name=F, italic=True)
ws.column_dimensions['A'].width = 28
for j in range(2,8): ws.column_dimensions[get_column_letter(j)].width = 12
# ---------------- Hourly validation ----------------
V = hr['val'][['case','window','method','n_hours','RMSE_cm','MAE_cm','R2']]
sheet('Hourly_Validation', V, widths={'window':26,'method':18}, numfmt={'RMSE_cm':'0.00','MAE_cm':'0.00','R2':'0.000'})
ws = wb.create_sheet('Hourly_Val_Summary'); NV = len(V)+1
ws.append(['method','mean RMSE (cm)','mean MAE (cm)','mean R2'])
for c in ws[1]: c.font = HF; c.fill = HFILL
for i, m in enumerate(['LI (hourly)','Tide + Z0','Tide + St','Tide + GP','Tide + GP-2M','Tide + GP-2M-SA'], 2):
    ws.cell(row=i, column=1, value=m)
    for j, col in ((2,'E'),(3,'F'),(4,'G')):
        c = ws.cell(row=i, column=j, value=f'=AVERAGEIFS(Hourly_Validation!${col}$2:${col}${NV},Hourly_Validation!$C$2:$C${NV},$A{i})'); c.number_format = '0.000' if j==4 else '0.00'
ws.cell(row=9, column=1, value='Ten withheld windows (Jun–Sep 2020, 2022–2025; Nov–Jan 2021/22–2025/26). Tide re-estimated without each window.').font = Font(name=F, italic=True)
ws.column_dimensions['A'].width = 22
for L in 'BCD': ws.column_dimensions[L].width = 16
# ---------------- Tides & gaps ----------------
tt = hr['tide_table'].sort_values('amplitude_m', ascending=False)
ws = sheet('Tidal_Constituents', tt, numfmt={'speed_deg_per_h':'0.0000000','amplitude_m':'0.0000','phase_deg_rel2020':'0.0'})
ws.cell(row=len(tt)+3, column=1, value=f'Z0 (mean level of hourly fit) = {hr["Z0"]:.4f} m. Phases are relative to 2020-01-01 00:00 without astronomical arguments (not Greenwich phases).').font = Font(name=F, italic=True)
sheet('Daily_Gaps', gaps, numfmt={})
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F: c.font = Font(name=F, bold=c.font.bold if c.font else False, italic=c.font.italic if c.font else False, color=c.font.color if c.font else None, size=c.font.size if c.font else 11)
wb.save('BangPakong_GP_Reconstruction.xlsx'); print('saved')
